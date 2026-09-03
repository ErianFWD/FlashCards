from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

CAFE = "4A2C20"
MIEL = "B77A22"
BEIGE = "F6EBDD"
CREMA = "FFFDF8"
AMARILLO = "E7B62D"
ROJO = "AD2C23"
GRIS = "6F625A"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def configure_document(doc, title, landscape=False):
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = Cm(1.4)
        section.right_margin = Cm(1.4)
    else:
        section.left_margin = Cm(2.1)
        section.right_margin = Cm(2.1)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(CAFE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 25, CAFE),
        ("Heading 1", 18, CAFE),
        ("Heading 2", 14, MIEL),
        ("Heading 3", 11.5, CAFE),
    ):
        style = doc.styles[name]
        style.font.name = "Georgia"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = f"FLASHCARD  ·  {title.upper()}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(GRIS)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(ROJO)
        else:
            paragraph.add_run(part)


def parse_table(lines, doc):
    rows = []
    for line in lines:
        values = [value.strip() for value in line.strip().strip("|").split("|")]
        rows.append(values)
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", value) for value in rows[1]):
        rows.pop(1)

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
            if row_index == 0:
                shade(cell, AMARILLO)
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(CAFE)
            elif row_index % 2 == 0:
                shade(cell, BEIGE)
    doc.add_paragraph()


def markdown_to_docx(source, output, title, landscape=False):
    doc = Document()
    configure_document(doc, title, landscape)
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False

    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            shade_paragraph = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), BEIGE)
            shade_paragraph.append(shd)
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            parse_table(table_lines, doc)
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, line[2:])
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:])
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:])
        elif re.match(r"^\d+\. ", line):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        else:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, line)
        index += 1

    properties = doc.core_properties
    properties.title = title
    properties.subject = "Práctica de Vibe Coding: solución de negocio"
    properties.author = "Proyecto Flashcard"
    doc.save(output)


markdown_to_docx(
    DOCS / "Bitacora_de_prompts.md",
    DOCS / "Bitacora_de_prompts.docx",
    "Bitácora de prompts",
    landscape=True,
)
markdown_to_docx(
    DOCS / "Reflexion_Vibe_Coding.md",
    DOCS / "Reflexion_Vibe_Coding.docx",
    "Reflexión sobre Vibe Coding",
    landscape=False,
)

print("Documentos creados correctamente.")
